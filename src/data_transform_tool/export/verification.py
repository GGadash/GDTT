"""Reopen exported artifacts and verify their critical data and structure.

Copyright (c) 2026 Akila DJ +. AI-assisted development: OpenAI Codex.
"""

from __future__ import annotations

import csv
import hashlib
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from data_transform_tool.datetime.custom_formats import decode
from data_transform_tool.domain.batches import TabularData, iter_rows
from data_transform_tool.domain.table import CellValue
from data_transform_tool.export.models import (
    ExportArtifact,
    ExportPlan,
    VerificationCheck,
    VerificationResult,
    VerificationStatus,
)
from data_transform_tool.export.naming import data_output_path
from data_transform_tool.export.writers import csv_output_row, output_row
from data_transform_tool.io.cancellation import (
    CancellationToken,
    InspectionCancelled,
)
from data_transform_tool.transformation.recipe import OutputFormat
from data_transform_tool.validation.serialization import serialize_delimited_row


def verify_artifact(
    artifact: ExportArtifact,
    expected: TabularData,
    plan: ExportPlan,
    cancellation: CancellationToken | None = None,
) -> VerificationResult:
    """Return explicit checks even when reopening the artifact fails."""
    checks = [
        VerificationCheck(
            "File exists",
            artifact.path.is_file(),
            str(artifact.path),
        )
    ]
    try:
        output_format = OutputFormat(artifact.kind)
        expected_path = data_output_path(plan.destination, plan.base_name, output_format)
        checks.append(
            VerificationCheck(
                "Filename",
                artifact.path == expected_path,
                f"Expected {expected_path.name}; found {artifact.path.name}.",
            )
        )
        if output_format is OutputFormat.CSV:
            checks.extend(_verify_csv(artifact.path, expected, plan, cancellation))
        else:
            checks.extend(_verify_xlsx(artifact.path, expected, plan, output_format, cancellation))
    except InspectionCancelled:
        raise
    except Exception as error:
        checks.append(
            VerificationCheck(
                "Readable output",
                False,
                f"{type(error).__name__}: {error}",
            )
        )
    return VerificationResult(artifact.path, _status(checks), tuple(checks))


def verify_artifacts(
    artifacts: tuple[ExportArtifact, ...],
    expected: TabularData,
    plan: ExportPlan,
    cancellation: CancellationToken | None = None,
) -> tuple[VerificationResult, ...]:
    return tuple(verify_artifact(artifact, expected, plan, cancellation) for artifact in artifacts)


def _verify_csv(
    path: Path,
    expected: TabularData,
    plan: ExportPlan,
    cancellation: CancellationToken | None,
) -> list[VerificationCheck]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        header = tuple(next(reader, ()))
        row_count = 0
        first_row: tuple[str, ...] | None = None
        last_row: tuple[str, ...] | None = None
        for row_count, row in enumerate(reader, start=1):
            if row_count % 10_000 == 0:
                _cancel(cancellation)
            current = tuple(row)
            if first_row is None:
                first_row = current
            last_row = current
    _cancel(cancellation)
    checks = _structural_checks(header, row_count, expected)
    checks.append(
        VerificationCheck(
            "CSV values and null representation",
            _csv_payload_digest(path, cancellation)
            == _expected_csv_digest(expected, plan, cancellation),
            "Every serialized row matches, including blank nulls and quoted empty text.",
        )
    )
    checks.append(_timestamp_check(header, first_row, last_row))
    return checks


def _verify_xlsx(
    path: Path,
    expected: TabularData,
    plan: ExportPlan,
    output_format: OutputFormat,
    cancellation: CancellationToken | None,
) -> list[VerificationCheck]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Data"]
        rows = worksheet.iter_rows(values_only=False)
        header_values = tuple(cell.value for cell in next(rows, ()))
        header = tuple(str(value) if value is not None else "" for value in header_values)
        expected_rows = (output_row(row, plan) for row in iter_rows(expected))
        values_match = True
        formats_match = True
        expected_formats = {
            name: spec.pattern
            for name, profile in plan.number_profiles
            if (spec := decode(profile)) is not None
        }
        row_count = 0
        first_row: tuple[Any, ...] | None = None
        last_row: tuple[Any, ...] | None = None
        for row_count, actual in enumerate(rows, start=1):
            if row_count % 10_000 == 0:
                _cancel(cancellation)
            try:
                wanted = next(expected_rows)
            except StopIteration:
                values_match = False
                continue
            actual_row: tuple[Any, ...] = tuple(cell.value for cell in actual)
            for name, cell in zip(header, actual, strict=False):
                if (
                    name in expected_formats
                    and isinstance(cell.value, (int, float))
                    and not isinstance(cell.value, bool)
                ):
                    formats_match = formats_match and cell.number_format == expected_formats[name]
            if first_row is None:
                first_row = actual_row
            last_row = actual_row
            if not _xlsx_row_equal(actual_row, wanted):
                values_match = False
        try:
            next(expected_rows)
        except StopIteration:
            pass
        else:
            values_match = False
        _cancel(cancellation)
        actual_count = int(worksheet.max_row or 0) - (1 if header else 0)
        checks = _structural_checks(header, max(actual_count, 0), expected)
        checks.append(
            VerificationCheck(
                "XLSX values and blank cells",
                values_match,
                "Every cell matches the expected transformed value; canonical nulls remain blank.",
            )
        )
        checks.append(_timestamp_check(header, first_row, last_row))
        if output_format is OutputFormat.XLSX_FORMATTED:
            checks.extend(_xlsx_layout_checks(path, plan))
        if expected_formats:
            checks.append(
                VerificationCheck(
                    "Column number formats",
                    formats_match,
                    "Numeric cells retain the selected Excel display masks.",
                )
            )
        return checks
    finally:
        workbook.close()


def _structural_checks(
    header: tuple[str, ...],
    row_count: int,
    expected: TabularData,
) -> list[VerificationCheck]:
    duplicate_index = sum(name.casefold() == "index" for name in header) > 1
    return [
        VerificationCheck(
            "Column names and order",
            header == expected.columns,
            f"Expected {expected.column_count} ordered columns; found {len(header)}.",
        ),
        VerificationCheck(
            "Row count",
            row_count == expected.row_count,
            f"Expected {expected.row_count:,}; found {row_count:,}.",
        ),
        VerificationCheck(
            "No duplicate Index column",
            not duplicate_index,
            "No accidental duplicate Index column was found."
            if not duplicate_index
            else "More than one column is named Index.",
        ),
    ]


def _xlsx_row_equal(actual: tuple[Any, ...], expected: tuple[CellValue, ...]) -> bool:
    if len(actual) != len(expected):
        return False
    return all(
        _normalize_xlsx_value(actual_value) == _normalize_xlsx_value(expected_value)
        for actual_value, expected_value in zip(actual, expected, strict=True)
    )


def _normalize_xlsx_value(value: object) -> object:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None).isoformat(sep=" ")
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).isoformat(sep=" ")
    if isinstance(value, Decimal):
        return float(value)
    return value


def _xlsx_layout_checks(path: Path, plan: ExportPlan) -> tuple[VerificationCheck, ...]:
    with zipfile.ZipFile(path) as archive:
        worksheet_xml = archive.read("xl/worksheets/sheet1.xml")
    has_filter = b"<autoFilter" in worksheet_xml
    has_freeze = b"<pane" in worksheet_xml and b'state="frozen"' in worksheet_xml
    return (
        VerificationCheck(
            "Formatted worksheet filter",
            has_filter if plan.xlsx_style.autofilter else not has_filter,
            f"AutoFilter is {'on' if has_filter else 'off'}.",
        ),
        VerificationCheck(
            "Formatted worksheet freeze",
            has_freeze if plan.xlsx_style.freeze_header else not has_freeze,
            f"Header freeze is {'on' if has_freeze else 'off'}.",
        ),
    )


def _timestamp_check(
    header: tuple[str, ...],
    first_row: tuple[Any, ...] | None,
    last_row: tuple[Any, ...] | None,
) -> VerificationCheck:
    candidates = [
        index
        for index, name in enumerate(header)
        if "time" in name.casefold() or name.casefold() in {"date", "period_start"}
    ]
    if not candidates or first_row is None or last_row is None:
        return VerificationCheck(
            "First and last timestamp",
            True,
            "No timestamp-like output column required a boundary check.",
            warning=True,
        )
    index = candidates[0]
    first = first_row[index] if index < len(first_row) else None
    last = last_row[index] if index < len(last_row) else None
    passed = _parseable_timestamp(first) and _parseable_timestamp(last)
    return VerificationCheck(
        "First and last timestamp",
        passed,
        f"First={first!s}; last={last!s}.",
    )


def _parseable_timestamp(value: object) -> bool:
    if isinstance(value, date | datetime):
        return True
    if value is None or not str(value).strip():
        return False
    try:
        datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
    except ValueError:
        return False
    return True


def _csv_payload_digest(
    path: Path,
    cancellation: CancellationToken | None,
) -> bytes:
    digest = hashlib.sha256()
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        while chunk := stream.read(1024 * 1024):
            _cancel(cancellation)
            digest.update(chunk.encode("utf-8"))
    return digest.digest()


def _expected_csv_digest(
    expected: TabularData,
    plan: ExportPlan,
    cancellation: CancellationToken | None,
) -> bytes:
    digest = hashlib.sha256()
    digest.update((serialize_delimited_row(expected.columns) + "\n").encode())
    for index, row in enumerate(iter_rows(expected), start=1):
        if index % 10_000 == 0:
            _cancel(cancellation)
        digest.update(
            (serialize_delimited_row(csv_output_row(row, plan, expected.columns)) + "\n").encode()
        )
    _cancel(cancellation)
    return digest.digest()


def _cancel(cancellation: CancellationToken | None) -> None:
    if cancellation is not None:
        cancellation.raise_if_cancelled()


def _status(checks: list[VerificationCheck]) -> VerificationStatus:
    if any(not check.passed and not check.warning for check in checks):
        return VerificationStatus.FAILED
    if any(not check.passed or check.warning for check in checks):
        return VerificationStatus.PASSED_WITH_WARNINGS
    return VerificationStatus.PASSED
