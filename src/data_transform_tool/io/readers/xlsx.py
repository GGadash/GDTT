"""Read-only streaming XLSX inspection and worksheet discovery.

Copyright (c) 2026 Akila DJ +. AI-assisted development: OpenAI Codex.
"""

from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from data_transform_tool.domain.errors import AppError
from data_transform_tool.io.cancellation import CancellationToken
from data_transform_tool.io.models import FileInspection, FileKind
from data_transform_tool.io.options import InspectionOptions
from data_transform_tool.io.reader import ProgressCallback
from data_transform_tool.io.readers.common import Row, inspect_row_stream, row_is_empty


class XlsxReader:
    """Inspect one selected worksheet using openpyxl read-only mode."""

    extensions = frozenset({".xlsx"})

    def available_worksheets(self, path: Path) -> tuple[str, ...]:
        workbook = self._open(path)
        try:
            return tuple(workbook.sheetnames)
        finally:
            workbook.close()

    def inspect(
        self,
        path: Path,
        options: InspectionOptions,
        cancellation: CancellationToken,
        progress: ProgressCallback | None = None,
    ) -> FileInspection:
        worksheets = self.available_worksheets(path)
        if not worksheets:
            raise AppError("The workbook does not contain any worksheets.")
        worksheet = options.worksheet or worksheets[0]
        if worksheet not in worksheets:
            raise AppError(f"Worksheet '{worksheet}' was not found in the workbook.")

        probe = self._raw_rows(path, worksheet)
        first_row = next(probe, None)
        second_row = next(probe, None)
        if first_row is None:
            raise AppError(f"Worksheet '{worksheet}' does not contain tabular rows.")

        if options.has_header is None:
            has_header = _looks_like_header(first_row, second_row)
            header_confidence = 0.8 if has_header else 0.6
        else:
            has_header = options.has_header
            header_confidence = 1.0

        initial_headers = (
            list(first_row)
            if has_header
            else [f"Column {index}" for index in range(1, len(first_row) + 1)]
        )

        def data_rows() -> Iterator[Row]:
            rows = self._raw_rows(path, worksheet)
            if has_header:
                next(rows, None)
            yield from rows

        warnings: list[str] = []
        if not has_header:
            warnings.append(
                "No header row was confidently detected; generated column names are provisional."
            )
        return inspect_row_stream(
            path=path,
            file_kind=FileKind.XLSX,
            initial_headers=initial_headers,
            rows_factory=data_rows,
            sample_limit=options.sample_limit,
            processing_strategy="openpyxl read-only streaming (two-pass preview)",
            estimated_memory_multiplier=12.0,
            worksheet=worksheet,
            available_worksheets=worksheets,
            encoding=None,
            encoding_confidence=None,
            delimiter=None,
            quote_character=None,
            header_detected=has_header,
            header_confidence=header_confidence,
            cancellation=cancellation,
            progress=progress,
            warnings=warnings,
        )

    def _raw_rows(self, path: Path, worksheet: str) -> Iterator[Row]:
        workbook = self._open(path)
        try:
            sheet = workbook[worksheet]
            for values in sheet.iter_rows(values_only=True):
                normalized = tuple(values)
                if not row_is_empty(normalized):
                    yield normalized
        finally:
            workbook.close()

    @staticmethod
    def _open(path: Path) -> Any:
        try:
            return load_workbook(path, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
            raise AppError(
                "The selected XLSX workbook could not be opened.", detail=str(error)
            ) from error


def _looks_like_header(first_row: Row, second_row: Row | None) -> bool:
    non_empty = [value for value in first_row if value is not None and str(value).strip()]
    if not non_empty:
        return False
    all_text = all(isinstance(value, str) for value in non_empty)
    unique_text = len({str(value).strip().casefold() for value in non_empty}) == len(non_empty)
    contains_names = all(
        any(character.isalpha() for character in str(value)) for value in non_empty
    )
    if not (all_text and unique_text and contains_names):
        return False
    if second_row is None:
        return True
    return any(isinstance(value, (int, float, bool, date, datetime, time)) for value in second_row)
