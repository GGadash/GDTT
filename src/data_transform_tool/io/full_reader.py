"""Guarded full-file readers for Phase 8 execution.

Copyright (c) 2026 Akila DJ +. AI-assisted development: OpenAI Codex.
"""

from __future__ import annotations

import csv
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal

from openpyxl import load_workbook

from data_transform_tool.domain.errors import AppError
from data_transform_tool.domain.spill import SpillTable, SpillWorkspace
from data_transform_tool.domain.table import CellValue, DataRow, DataTable
from data_transform_tool.io.cancellation import CancellationToken
from data_transform_tool.io.models import FileInspection, FileKind

ProgressCallback = Callable[[str, int | None], None]


@dataclass(frozen=True)
class FullReadResult:
    table: DataTable
    strategy: str
    source_bytes: int


@dataclass(frozen=True)
class FullSpillReadResult:
    table: SpillTable
    strategy: str
    source_bytes: int


def read_full_table(
    inspection: FileInspection,
    cancellation: CancellationToken,
    *,
    progress: ProgressCallback | None = None,
    max_estimated_memory_bytes: int = 1_000_000_000,
) -> FullReadResult:
    """Read every row after a visible memory guard; never reuse bounded preview slices."""
    if inspection.estimated_memory_bytes > max_estimated_memory_bytes:
        raise AppError(
            "The estimated full-file working set exceeds the safe Phase 8 limit.",
            detail=(
                f"Estimated {inspection.estimated_memory_bytes:,} bytes; limit "
                f"{max_estimated_memory_bytes:,}. Export CSV subsets or wait for the Phase 9 "
                "benchmarked out-of-core backend."
            ),
        )
    rows = tuple(iter_full_rows(inspection, cancellation, progress=progress))
    if progress is not None:
        progress("Source loaded", len(rows))
    return FullReadResult(
        DataTable(inspection.column_names, tuple(rows)),
        "streaming source read with guarded semantic materialization",
        inspection.file_size_bytes,
    )


def read_full_spill(
    inspection: FileInspection,
    cancellation: CancellationToken,
    workspace: SpillWorkspace,
    *,
    progress: ProgressCallback | None = None,
) -> FullSpillReadResult:
    """Stream the complete source into a private replayable spill table."""
    table = workspace.write_table(
        "input",
        inspection.column_names,
        iter_full_rows(inspection, cancellation, progress=progress),
        cancellation,
    )
    if progress is not None:
        progress("Source staged", table.row_count)
    return FullSpillReadResult(
        table,
        "streaming source read into private bounded-memory spill batches",
        inspection.file_size_bytes,
    )


def iter_full_rows(
    inspection: FileInspection,
    cancellation: CancellationToken,
    *,
    progress: ProgressCallback | None = None,
) -> Iterator[DataRow]:
    """Yield normalized complete-source rows without using inspection previews."""
    expected = inspection.column_count
    for row_number, row in enumerate(_rows(inspection), start=1):
        cancellation.raise_if_cancelled()
        normalized = tuple(_cell(value) for value in row[:expected])
        if len(normalized) < expected:
            normalized = (*normalized, *(None for _ in range(expected - len(normalized))))
        if progress is not None and row_number % 10_000 == 0:
            progress("Reading complete source", row_number)
        yield normalized


def _rows(inspection: FileInspection) -> Iterator[tuple[object | None, ...]]:
    if inspection.file_kind is FileKind.XLSX:
        yield from _xlsx_rows(inspection)
    else:
        yield from _delimited_rows(inspection)


def _delimited_rows(inspection: FileInspection) -> Iterator[tuple[object | None, ...]]:
    encoding = inspection.encoding or "utf-8"
    delimiter = inspection.delimiter or ","
    quote = inspection.quote_character or '"'
    with inspection.path.open("r", encoding=encoding, newline="") as source:
        reader = csv.reader(source, delimiter=delimiter, quotechar=quote)
        if inspection.header_detected:
            next(reader, None)
        for row in reader:
            if any(value != "" for value in row):
                yield tuple(row)


def _xlsx_rows(inspection: FileInspection) -> Iterator[tuple[object | None, ...]]:
    if inspection.worksheet is None:
        raise AppError("The XLSX execution source requires a selected worksheet.")
    workbook = load_workbook(inspection.path, read_only=True, data_only=True)
    try:
        sheet = workbook[inspection.worksheet]
        rows = sheet.iter_rows(values_only=True)
        if inspection.header_detected:
            next(rows, None)
        for row in rows:
            if any(value is not None and str(value).strip() for value in row):
                yield tuple(row)
    finally:
        workbook.close()


def _cell(value: object | None) -> CellValue:
    if value is None or isinstance(
        value, str | int | float | bool | Decimal | date | datetime | time | timedelta
    ):
        return value
    return str(value)
