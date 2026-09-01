"""Phase 2 reader and profiling tests using only synthetic data."""

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from data_transform_tool.domain.errors import AppError
from data_transform_tool.io import FileInspector, FileKind, SemanticType
from data_transform_tool.io.cancellation import CancellationToken, InspectionCancelled
from data_transform_tool.io.options import InspectionOptions
from data_transform_tool.io.readers import common

FIXTURES = Path(__file__).parents[1] / "fixtures"


def test_csv_profile_is_complete_and_does_not_modify_source() -> None:
    path = FIXTURES / "monitoring_gap.csv"
    original = path.read_bytes()

    result = FileInspector.default().inspect(path)

    assert path.read_bytes() == original
    assert result.file_kind is FileKind.CSV
    assert result.row_count == 10
    assert result.column_count == 5
    assert result.column_names == (
        "DateTime_UTC",
        "Source",
        "PM2.5",
        "Temperature",
        "Unused",
    )
    assert result.encoding == "utf-8"
    assert result.delimiter == ","
    assert result.header_detected is True
    assert result.empty_columns == ("Unused",)
    assert {marker.value: marker.count for marker in result.potential_missing_markers} == {
        "N/A": 1,
        "-999": 1,
    }

    profiles = {column.name: column for column in result.columns}
    assert profiles["DateTime_UTC"].inferred_type is SemanticType.DATETIME
    assert profiles["Source"].inferred_type is SemanticType.CATEGORY
    assert profiles["PM2.5"].inferred_type is SemanticType.DECIMAL
    assert profiles["Temperature"].inferred_type is SemanticType.DECIMAL
    assert result.likely_datetime_column == "DateTime_UTC"
    assert result.likely_interval_seconds == 3600.0
    assert result.likely_interval_label == "1 hour"
    assert [preview.start_row for preview in result.previews] == [1, 3, 6]
    assert [len(preview.rows) for preview in result.previews] == [5, 5, 5]
    assert result.large_file is False
    assert "constant-memory" in result.processing_strategy


def test_tsv_supports_explicit_no_header_override(tmp_path: Path) -> None:
    path = tmp_path / "readings.tsv"
    path.write_text("alpha\t1\nbeta\t2\n", encoding="utf-8")

    result = FileInspector.default().inspect(path, InspectionOptions(has_header=False))

    assert result.file_kind is FileKind.TSV
    assert result.delimiter == "\t"
    assert result.header_detected is False
    assert result.column_names == ("Column 1", "Column 2")
    assert result.row_count == 2
    assert result.previews[0].rows[0] == ("alpha", "1")


def test_non_utf8_txt_encoding_is_detected_and_decoded(tmp_path: Path) -> None:
    path = tmp_path / "regional.txt"
    path.write_bytes(b"name;value\ncaf\xe9;1\n")

    result = FileInspector.default().inspect(
        path,
        InspectionOptions(delimiter=";", has_header=True),
    )

    assert result.encoding not in {None, "utf-8"}
    assert result.previews[0].rows[0][0] == "café"


def test_xlsx_lists_and_inspects_selected_worksheet(tmp_path: Path) -> None:
    path = tmp_path / "monitoring.xlsx"
    workbook = Workbook()
    monitoring = workbook.active
    monitoring.title = "Monitoring"
    monitoring.append(["Timestamp", "PM2.5", "Unused"])
    start = datetime(2025, 8, 20, 8)
    for index in range(4):
        monitoring.append([start + timedelta(hours=index), 10.5 + index, None])
    notes = workbook.create_sheet("Notes")
    notes.append(["Key", "Value"])
    notes.append(["Site", "Synthetic"])
    workbook.save(path)

    inspector = FileInspector.default()
    assert inspector.available_worksheets(path) == ("Monitoring", "Notes")

    result = inspector.inspect(path, InspectionOptions(worksheet="Monitoring"))

    assert result.file_kind is FileKind.XLSX
    assert result.worksheet == "Monitoring"
    assert result.available_worksheets == ("Monitoring", "Notes")
    assert result.row_count == 4
    assert result.column_names == ("Timestamp", "PM2.5", "Unused")
    assert result.encoding is None
    assert result.delimiter is None
    assert result.likely_datetime_column == "Timestamp"
    assert result.likely_interval_seconds == 3600.0
    assert result.empty_columns == ("Unused",)


def test_invalid_worksheet_and_unsupported_file_are_safe_errors(tmp_path: Path) -> None:
    workbook_path = tmp_path / "book.xlsx"
    workbook = Workbook()
    workbook.save(workbook_path)

    inspector = FileInspector.default()
    with pytest.raises(AppError, match="Worksheet 'Missing'"):
        inspector.inspect(workbook_path, InspectionOptions(worksheet="Missing"))

    unsupported = tmp_path / "values.json"
    unsupported.write_text("{}", encoding="utf-8")
    with pytest.raises(AppError, match="Unsupported file type"):
        inspector.inspect(unsupported)


def test_pre_cancelled_inspection_stops_without_mutation() -> None:
    path = FIXTURES / "monitoring_gap.csv"
    original = path.read_bytes()
    cancellation = CancellationToken()
    cancellation.cancel()

    with pytest.raises(InspectionCancelled):
        FileInspector.default().inspect(path, cancellation=cancellation)

    assert path.read_bytes() == original


def test_large_file_policy_reports_streaming_requirement(monkeypatch) -> None:  # type: ignore[no-untyped-def]
    monkeypatch.setattr(common, "LARGE_FILE_BYTES", 1)

    result = FileInspector.default().inspect(FIXTURES / "monitoring_gap.csv")

    assert result.large_file is True
    assert any("lazy/streaming" in warning for warning in result.warnings)


def test_ambiguous_regional_dates_require_confirmation(tmp_path: Path) -> None:
    path = tmp_path / "ambiguous.csv"
    path.write_text(
        "when,value\n01/02/2025,1\n02/03/2025,2\n03/04/2025,3\n",
        encoding="utf-8",
    )

    result = FileInspector.default().inspect(path, InspectionOptions(has_header=True))

    assert result.columns[0].inferred_type is SemanticType.DATE
    assert "Ambiguous regional date order" in result.columns[0].warnings[0]


@pytest.mark.parametrize("delimiter", ["", "::"])
def test_delimiter_override_must_be_one_character(delimiter: str) -> None:
    with pytest.raises(ValueError, match="exactly one"):
        InspectionOptions(delimiter=delimiter)
