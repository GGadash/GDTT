"""Bounded-memory tabular contract and deterministic spill cleanup coverage."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from data_transform_tool.domain.batches import iter_rows
from data_transform_tool.domain.spill import SpillWorkspace
from data_transform_tool.domain.table import DataRow, DataTable
from data_transform_tool.export import ExportPlan, VerificationStatus
from data_transform_tool.export.verification import verify_artifacts
from data_transform_tool.export.writers import write_data_outputs
from data_transform_tool.io.cancellation import CancellationToken, InspectionCancelled
from data_transform_tool.io.full_reader import read_full_spill
from data_transform_tool.io.inspector import FileInspector
from data_transform_tool.io.options import InspectionOptions
from data_transform_tool.transformation.recipe import OutputFormat


def _rows(count: int) -> tuple[DataRow, ...]:
    return tuple((index, None if index % 3 == 0 else f"value-{index}") for index in range(count))


def test_spill_table_replays_batches_slices_writes_and_verifies(tmp_path: Path) -> None:
    token = CancellationToken()
    with SpillWorkspace(parent=tmp_path) as workspace:
        workspace_path = workspace.path
        table = workspace.write_table("output", ("Index", "Value"), _rows(23), token, batch_size=7)

        assert [len(batch.rows) for batch in table.iter_batches(10)] == [10, 10, 3]
        assert table.read_rows(9, 4) == _rows(23)[9:13]
        assert tuple(iter_rows(table, 6)) == _rows(23)

        plan = ExportPlan(tmp_path / "exports", "batch", (OutputFormat.CSV,))
        artifacts = write_data_outputs(table, plan, token)
        verification = verify_artifacts(artifacts, table, plan, token)

        assert verification[0].status is VerificationStatus.PASSED_WITH_WARNINGS
        assert workspace_path.is_dir()
    assert not workspace_path.exists()


def test_cancelled_spill_write_removes_partial_database(tmp_path: Path) -> None:
    token = CancellationToken()
    workspace = SpillWorkspace(parent=tmp_path)

    def cancelling_rows() -> tuple[DataRow, ...]:
        token.cancel()
        return _rows(2)

    with pytest.raises(InspectionCancelled):
        workspace.write_table(
            "cancelled",
            ("Index", "Value"),
            cancelling_rows(),
            token,
            batch_size=1,
        )

    assert not tuple(workspace.path.glob("*.sqlite3"))
    workspace.close()
    assert not workspace.path.exists()


def test_cancelled_batch_export_removes_same_directory_temporary_file(
    tmp_path: Path,
) -> None:
    class CancelDuringWrite(CancellationToken):
        def __init__(self) -> None:
            super().__init__()
            self.checks = 0

        def raise_if_cancelled(self) -> None:
            self.checks += 1
            if self.checks >= 3:
                self.cancel()
            super().raise_if_cancelled()

    setup_token = CancellationToken()
    with SpillWorkspace(parent=tmp_path) as workspace:
        table = workspace.write_table(
            "large",
            ("Index", "Value"),
            _rows(25_000),
            setup_token,
        )
        destination = tmp_path / "cancelled-export"
        plan = ExportPlan(destination, "cancelled", (OutputFormat.CSV,))

        with pytest.raises(InspectionCancelled):
            write_data_outputs(table, plan, CancelDuringWrite())

        assert not (destination / "cancelled.csv").exists()
        assert not tuple(destination.glob(".cancelled-*.csv"))


def test_cancelled_stream_verification_propagates_cancellation(tmp_path: Path) -> None:
    table = DataTable(("Timestamp",), (("2026-01-01T00:00:00",),))
    plan = ExportPlan(tmp_path, "verify-cancel", (OutputFormat.CSV,))
    artifact = write_data_outputs(table, plan, CancellationToken())[0]
    cancelled = CancellationToken()
    cancelled.cancel()

    with pytest.raises(InspectionCancelled):
        verify_artifacts((artifact,), table, plan, cancelled)


def test_spill_reader_preserves_non_utf8_delimited_fallback(tmp_path: Path) -> None:
    source = tmp_path / "regional.txt"
    source.write_bytes(b"name;value\ncaf\xe9;1\n")
    inspection = FileInspector.default().inspect(
        source,
        InspectionOptions(delimiter=";", has_header=True),
    )
    token = CancellationToken()

    with SpillWorkspace(parent=tmp_path) as workspace:
        result = read_full_spill(inspection, token, workspace)

        assert result.table.read_rows(0, 1) == (("café", "1"),)


def test_spill_reader_preserves_read_only_xlsx_fallback(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monitoring"
    sheet.append(["Timestamp", "Value"])
    sheet.append(["2026-01-01T00:00:00", 1.5])
    workbook.save(source)
    inspection = FileInspector.default().inspect(
        source,
        InspectionOptions(worksheet="Monitoring"),
    )
    token = CancellationToken()

    with SpillWorkspace(parent=tmp_path) as workspace:
        result = read_full_spill(inspection, token, workspace)

        assert result.table.read_rows(0, 1) == (("2026-01-01T00:00:00", 1.5),)
