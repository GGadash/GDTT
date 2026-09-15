"""0.10.0 explicit locale, precision and inert custom-format contracts."""

from dataclasses import replace
from datetime import datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from data_transform_tool.app.export_workflow import execute_export, prepare_export
from data_transform_tool.app.reformat_configuration import (
    build_transformation_recipe,
    create_reformat_draft,
)
from data_transform_tool.app.reformat_preview import build_proposed_preview
from data_transform_tool.app.template_application import apply_reformat_template
from data_transform_tool.datetime.custom_formats import FieldFormat, apply_profile, decode
from data_transform_tool.domain.table import DataTable
from data_transform_tool.export import ExportPlan, VerificationStatus
from data_transform_tool.io.cancellation import CancellationToken
from data_transform_tool.io.inspector import FileInspector
from data_transform_tool.io.models import SemanticType
from data_transform_tool.io.options import InspectionOptions
from data_transform_tool.timezone.zones import resolve_zone
from data_transform_tool.transformation.base import InvalidValuePolicy, TransformationError
from data_transform_tool.transformation.recipe import OutputFormat


@pytest.mark.parametrize(
    ("mask", "separator", "source", "expected"),
    [
        ("0.00", ".", "1.2345", "1.2345"),
        ("0.00", ",", "1,2345", "1.2345"),
        ("#,##0.00", ",", "1.234,567", "1234.567"),
        ("#,##0.00", ",", "1\u202f234,567", "1234.567"),
        ("#,##0.00", ".", "1,234.567", "1234.567"),
    ],
)
def test_input_locale_never_rounds(mask, separator, source, expected):
    assert FieldFormat("number", mask, separator).number(source, parsing=True) == Decimal(expected)


def test_iso_compact_offset_to_utc_preview_recipe_and_full_exports(tmp_path):
    path = tmp_path / "iso.csv"
    path.write_text("When,Value\n20260101T010203+0530,1.23\n", encoding="utf-8")
    inspection = FileInspector.default().inspect(path)
    draft = replace(create_reformat_draft(inspection), gap_fill_enabled=False)
    draft = draft.update_column(
        replace(
            draft.column("When"),
            output_type=SemanticType.DATETIME,
            input_profile="iso_basic_offset",
            output_profile="iso_utc",
        )
    )
    preview = build_proposed_preview(inspection, draft)
    assert preview.error is None
    expected = "2025-12-31T19:32:03Z"
    assert preview.rows[0][0] == expected
    recipe = build_transformation_recipe(draft)
    restored = apply_reformat_template(create_reformat_draft(inspection), recipe)
    assert restored.column("When").input_profile == "iso_basic_offset"
    assert restored.column("When").output_profile == "iso_utc"
    prepared = prepare_export(inspection, draft, recipe, CancellationToken())
    try:
        result = execute_export(
            prepared, ExportPlan(tmp_path, "result", tuple(OutputFormat)), CancellationToken()
        )
        assert all(v.status is not VerificationStatus.FAILED for v in result.verifications)
        assert all(check.passed for v in result.verifications for check in v.checks)
        assert expected in (tmp_path / "result.csv").read_text(encoding="utf-8-sig")
        for name in ("result_P.xlsx", "result_F.xlsx"):
            workbook = load_workbook(tmp_path / name)
            try:
                assert workbook.active["A2"].value == expected
            finally:
                workbook.close()
    finally:
        prepared.close()


@pytest.mark.parametrize("source", ["1.2345", "1,2,3", "1 234,56"])
def test_input_separator_does_not_guess(source):
    with pytest.raises(ValueError):
        FieldFormat("number", "0.00", ",").number(source, parsing=True)


def test_numeric_masks_round_or_preserve_explicitly():
    rounded = FieldFormat("number", "0.00")
    assert rounded.number(Decimal("1.2345")) == Decimal("1.23")
    assert rounded.number(Decimal("-1.235")) == Decimal("-1.24")
    assert rounded.display(1) == "1.00"
    preserved = replace(rounded, preserve=True)
    assert preserved.number(Decimal("1.2345")) == Decimal("1.2345")
    assert preserved.display(Decimal("1.2345")) == "1.23"
    assert preserved.display(Decimal("1.2345"), csv=True) == "1.2345"
    assert FieldFormat("number", "#,##0.00", ",").display(Decimal("1234.567")) == "1.234,57"
    assert FieldFormat("number", "0.###").display(Decimal("1.200")) == "1.2"
    assert decode(preserved.encode()) == preserved


def test_custom_temporal_and_text_boolean_roundtrip():
    source = DataTable(("v",), (("2026-09-11 13:45:27",), (None,)))
    profile = FieldFormat("datetime", "yyyy-MM-dd HH:mm:ss").encode()
    parsed = apply_profile(source, "v", profile, InvalidValuePolicy.STOP, parsing=True).table
    assert parsed.rows[0][0] == datetime(2026, 9, 11, 13, 45, 27)
    assert (
        apply_profile(parsed, "v", profile, InvalidValuePolicy.STOP, parsing=False).table == source
    )
    for kind, mask, value in (
        ("text", "pre-{value}-post", "reading"),
        ("boolean", "Oui|Non", True),
    ):
        original = DataTable(("v",), ((value,), (None,)))
        encoded = FieldFormat(kind, mask).encode()
        formatted = apply_profile(
            original, "v", encoded, InvalidValuePolicy.STOP, parsing=False
        ).table
        assert (
            apply_profile(formatted, "v", encoded, InvalidValuePolicy.STOP, parsing=True).table
            == original
        )


def test_invalid_custom_values_obey_null_and_stop():
    table = DataTable(("v",), (("wrong",), (None,)))
    profile = FieldFormat("number", "0.00", ",").encode()
    result = apply_profile(table, "v", profile, InvalidValuePolicy.NULL, parsing=True)
    assert result.table.rows == ((None,), (None,))
    assert result.diagnostics[0].count == 1
    with pytest.raises(TransformationError):
        apply_profile(table, "v", profile, InvalidValuePolicy.STOP, parsing=True)


@pytest.mark.parametrize("preserve", [False, True])
def test_locale_precision_end_to_end_and_saved_recipe(tmp_path, preserve):
    path = tmp_path / "sample.csv"
    path.write_text('Value,Station\n"1,2345",A\n"2,3456",B\n', encoding="utf-8")
    inspection = FileInspector.default().inspect(
        path, InspectionOptions(), cancellation=CancellationToken()
    )
    draft = replace(create_reformat_draft(inspection), gap_fill_enabled=False)
    draft = draft.update_column(
        replace(
            draft.column("Value"),
            output_type=SemanticType.DECIMAL,
            input_profile=FieldFormat("number", "0.00", ",").encode(),
            output_profile=FieldFormat("number", "0.00", ",", preserve).encode(),
        )
    )
    preview = build_proposed_preview(inspection, draft)
    assert preview.error is None
    assert preview.rows[0][0] == "1,23"
    recipe = build_transformation_recipe(draft)
    restored = apply_reformat_template(create_reformat_draft(inspection), recipe)
    assert restored.column("Value").output_profile == draft.column("Value").output_profile
    prepared = prepare_export(inspection, draft, recipe, CancellationToken())
    try:
        result = execute_export(
            prepared, ExportPlan(tmp_path, "result", tuple(OutputFormat)), CancellationToken()
        )
        assert all(v.status is not VerificationStatus.FAILED for v in result.verifications)
        assert all(check.passed for v in result.verifications for check in v.checks)
        assert ('"1,2345"' if preserve else '"1,23"') in (tmp_path / "result.csv").read_text(
            encoding="utf-8-sig"
        )
        for name in ("result_P.xlsx", "result_F.xlsx"):
            workbook = load_workbook(tmp_path / name)
            try:
                assert workbook.active["A2"].value == (1.2345 if preserve else 1.23)
                assert workbook.active["A2"].number_format == "0.00"
            finally:
                workbook.close()
    finally:
        prepared.close()


def test_fixed_offset_and_named_zone_choices():
    for zone in ("Asia/Colombo", "+05:30", "+5.5", "Colombo"):
        assert datetime(2026, 9, 11, tzinfo=resolve_zone(zone)).utcoffset() == timedelta(hours=5.5)
    assert datetime(2026, 9, 11, tzinfo=resolve_zone("+05:45")).utcoffset() == timedelta(hours=5.75)
    with pytest.raises(ValueError):
        resolve_zone("+25:00")
